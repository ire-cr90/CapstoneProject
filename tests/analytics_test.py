"""Tests for the pandas ``DataProcessor`` and the ``ReportGenerator``."""

from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.db.manager import DatabaseManager
from src.exceptions import ReportGenerationError
from src.processing.analytics import DataProcessor
from src.reporting.generator import ReportGenerator


@pytest.fixture
def works() -> pd.DataFrame:
    """Four works with missing years, languages and edition counts."""
    return pd.DataFrame(
        [
            {
                "work_key": "OL1W",
                "title": "Dune",
                "first_publish_year": 1965,
                "languages": '["eng", "spa"]',
                "isbns": ["9780441172719"],
                "subjects": ["Science fiction", "Deserts"],
                "edition_count": 155,
                "source": "api",
                "created_at": "2026-01-01T12:00:00+00:00",
                "updated_at": "2026-01-02T12:00:00+00:00",
            },
            {
                "work_key": "OL2W",
                "title": "Dune Messiah",
                "first_publish_year": "1969",
                "languages": ["eng"],
                "isbns": [],
                "subjects": ["science fiction "],
                "edition_count": 45,
                "source": "api",
                "created_at": "2026-01-01T12:00:00+00:00",
                "updated_at": "not a date",
            },
            {
                "work_key": "OL3W",
                "title": "Cien años de soledad",
                "first_publish_year": "unknown",
                "languages": [],
                "isbns": ["9780307474728"],
                "subjects": [],
                "edition_count": None,
                "source": "catalog",
                "created_at": None,
                "updated_at": None,
            },
            {
                "work_key": "OL4W",
                "title": "Children of Dune",
                "first_publish_year": 1965,
                "languages": None,
                "isbns": None,
                "subjects": None,
                "edition_count": "30",
                "source": "api",
                "created_at": None,
                "updated_at": None,
            },
        ]
    )


@pytest.fixture
def authors() -> pd.DataFrame:
    """Two authors."""
    return pd.DataFrame(
        [
            {"author_key": "OL1A", "name": "Frank Herbert"},
            {"author_key": "OL2A", "name": "Gabriel García Márquez"},
        ]
    )


@pytest.fixture
def links() -> pd.DataFrame:
    """Author links: OL4W has no author."""
    return pd.DataFrame(
        [
            {"work_key": "OL1W", "author_key": "OL1A"},
            {"work_key": "OL2W", "author_key": "OL1A"},
            {"work_key": "OL3W", "author_key": "OL2A"},
        ]
    )


@pytest.fixture
def benchmarks() -> pd.DataFrame:
    """Two benchmark runs."""
    return pd.DataFrame(
        [
            {
                "subject": "art",
                "pages": 2,
                "sync_seconds": 4.0,
                "async_seconds": 1.0,
                "speedup": 4.0,
            },
            {
                "subject": "art",
                "pages": 2,
                "sync_seconds": 2.0,
                "async_seconds": 1.0,
                "speedup": 2.0,
            },
        ]
    )


@pytest.fixture
def processor(works, authors, links, benchmarks) -> DataProcessor:
    """A processor built from the fake DataFrames."""
    return DataProcessor(works, authors, links, benchmarks, cache_hit_ratio=0.75)


def test_missing_values_are_flagged_and_reported(processor: DataProcessor) -> None:
    """Missing years, languages, edition counts and authors are counted."""
    assert processor.missing_values_report() == {
        "first_publish_year": 1,
        "languages": 2,
        "edition_count": 1,
        "authors": 1,
    }
    filled = processor.works.set_index("work_key")["languages_filled"]
    assert filled["OL3W"] == ["und"]
    assert filled["OL1W"] == ["eng", "spa"]


def test_date_parsing_and_decades(processor: DataProcessor) -> None:
    """Years become nullable integers, dates and decades; timestamps are parsed as UTC."""
    works = processor.works.set_index("work_key")
    assert str(works["first_publish_year"].dtype) == "Int64"
    assert works.loc["OL2W", "first_publish_year"] == 1969
    assert pd.isna(works.loc["OL3W", "first_publish_year"])
    assert works.loc["OL1W", "publish_date"] == pd.Timestamp("1965-01-01")
    assert works.loc["OL2W", "decade"] == 1960
    assert works.loc["OL1W", "updated_at"] == pd.Timestamp("2026-01-02T12:00:00Z")
    assert pd.isna(works.loc["OL2W", "updated_at"])


def test_works_with_authors_join(processor: DataProcessor) -> None:
    """The join keeps works without authors and attaches author names."""
    joined = processor.works_with_authors()
    assert len(joined) == 4
    assert joined.set_index("work_key").loc["OL3W", "author_name"] == "Gabriel García Márquez"
    assert pd.isna(joined.set_index("work_key").loc["OL4W", "author_name"])


def test_groupings(processor: DataProcessor) -> None:
    """Top years, decades, languages, subjects and authors are aggregated correctly."""
    years = processor.top_publication_years()
    assert years.to_dict("records")[0] == {"year": 1965, "count": 2}
    assert processor.works_per_decade().to_dict("records") == [{"decade": 1960, "count": 3}]
    languages = dict(zip(*processor.top_languages().T.values, strict=True))
    assert languages == {"eng": 2, "und": 2, "spa": 1}
    subjects = processor.top_subjects(n=1).to_dict("records")
    assert subjects == [{"subject": "science fiction", "count": 2}]
    top_authors = processor.top_authors().to_dict("records")
    assert top_authors[0] == {"author_key": "OL1A", "author": "Frank Herbert", "works": 2}
    assert processor.average_edition_count() == pytest.approx(76.67)


def test_benchmark_summary(processor: DataProcessor) -> None:
    """The benchmark summary averages the speedups."""
    summary = processor.benchmark_summary()
    assert summary["runs"] == 2
    assert summary["mean_speedup"] == pytest.approx(3.0)
    assert summary["best_speedup"] == pytest.approx(4.0)
    assert summary["mean_sync_seconds"] == pytest.approx(3.0)


def test_kpis(processor: DataProcessor) -> None:
    """KPIs combine counts, years, cache ratio, speedup and completeness."""
    kpis = processor.kpis()
    assert kpis["total_works"] == 4
    assert kpis["total_authors"] == 2
    assert kpis["unique_languages"] == 2
    assert kpis["oldest_year"] == 1965
    assert kpis["newest_year"] == 1969
    assert kpis["cache_hit_ratio"] == 0.75
    assert kpis["average_async_speedup"] == pytest.approx(3.0)
    assert kpis["data_completeness_pct"] == pytest.approx(round(17 / 24 * 100, 2))


def test_works_table_and_export_records(processor: DataProcessor) -> None:
    """The flat table and JSON records contain authors and JSON-friendly values."""
    table = processor.works_table()
    assert list(table["title"]) == [
        "Children of Dune",
        "Cien años de soledad",
        "Dune",
        "Dune Messiah",
    ]
    assert table.set_index("work_key").loc["OL1W", "languages"] == "eng, spa"
    records = {record["work_key"]: record for record in processor.export_records()}
    assert records["OL1W"]["authors"] == ["Frank Herbert"]
    assert records["OL3W"]["first_publish_year"] is None
    assert records["OL4W"]["edition_count"] == 30
    json.dumps(list(records.values()))


def test_impossible_years_are_treated_as_missing() -> None:
    """Year 0, negative years and future years count as missing."""
    processor = DataProcessor(
        [
            {"work_key": "A", "title": "A", "first_publish_year": 0},
            {"work_key": "B", "title": "B", "first_publish_year": -50},
            {"work_key": "C", "title": "C", "first_publish_year": 9999},
            {"work_key": "D", "title": "D", "first_publish_year": 1605},
        ]
    )
    assert processor.missing_values_report()["first_publish_year"] == 3
    assert processor.kpis()["oldest_year"] == 1605


def test_empty_processor_is_safe() -> None:
    """An empty database produces empty results instead of errors."""
    processor = DataProcessor([])
    assert processor.top_publication_years().empty
    assert processor.top_languages().empty
    assert processor.top_subjects().empty
    assert processor.top_authors().empty
    assert processor.average_edition_count() is None
    assert processor.kpis()["total_works"] == 0
    assert processor.kpis()["data_completeness_pct"] == 0.0
    assert processor.export_records() == []


def test_from_database(tmp_path: Path) -> None:
    """The processor loads works, authors, links, benchmarks and cache stats from SQLite."""
    with DatabaseManager(tmp_path / "library.db") as db:
        db.upsert_works_batch(
            [
                {
                    "work_key": "OL1W",
                    "title": "Dune",
                    "first_publish_year": 1965,
                    "languages": ["eng"],
                    "authors": [{"author_key": "OL1A", "name": "Frank Herbert"}],
                }
            ]
        )
        db.record_benchmark("art", 1, 2.0, 1.0, 1)
        db.get_cached("missing")
        processor = DataProcessor.from_database(db)
    assert processor.kpis()["total_works"] == 1
    assert processor.top_authors().iloc[0]["author"] == "Frank Herbert"
    assert processor.benchmark_summary()["mean_speedup"] == pytest.approx(2.0)
    assert processor.cache_hit_ratio == 0.0


def fixed_clock() -> datetime:
    """Return a fixed date for predictable report file names."""
    return datetime(2026, 10, 16, 9, 30)


def test_report_generator_creates_three_files(processor: DataProcessor, tmp_path: Path) -> None:
    """Excel, CSV and JSON files are created in a new output folder with dated names."""
    output = tmp_path / "outputs" / "nested"
    paths = ReportGenerator(output, clock=fixed_clock).generate_all(processor)
    assert paths["excel"].name == "library_report_20261016.xlsx"
    assert paths["csv"].name == "library_summary_20261016.csv"
    assert paths["json"].name == "library_export_20261016.json"
    assert all(path.is_file() and path.parent == output for path in paths.values())


def test_excel_report_sheets_and_formatting(processor: DataProcessor, tmp_path: Path) -> None:
    """The workbook has every sheet, styled frozen headers, autofilters and a chart."""
    path = ReportGenerator(tmp_path, clock=fixed_clock).export_excel(processor)
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Summary",
        "Works",
        "Top Years",
        "Languages",
        "Authors",
        "Subjects",
        "Benchmarks",
    ]
    works = workbook["Works"]
    header = works["A1"]
    assert header.font.bold and header.font.color.rgb.endswith("FFFFFF")
    assert header.fill.fgColor.rgb.endswith("1F4E78")
    assert works.freeze_panes == "A2"
    assert works.auto_filter.ref == works.dimensions
    assert works.column_dimensions["B"].width > 8
    assert works.max_row == 5
    summary = {row[1]: row[2] for row in workbook["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["Total Works"] == 4
    assert workbook["Summary"]["C2"].alignment.horizontal == "right"
    assert len(workbook["Top Years"]._charts) == 1
    speedup_cell = workbook["Benchmarks"]["G2"]
    assert speedup_cell.value == 4.0
    assert speedup_cell.number_format == '0.00"x"'


def test_csv_summary_contents(processor: DataProcessor, tmp_path: Path) -> None:
    """The CSV summary lists years and languages with clear columns."""
    path = ReportGenerator(tmp_path, clock=fixed_clock).export_csv_summary(processor)
    with path.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert list(rows[0]) == ["category", "value", "count"]
    assert {"category": "publication_year", "value": "1965", "count": "2"} in rows
    assert {"category": "language", "value": "eng", "count": "2"} in rows


def test_json_export_is_minified(processor: DataProcessor, tmp_path: Path) -> None:
    """The JSON export is minified UTF-8 with metadata and records."""
    path = ReportGenerator(tmp_path, clock=fixed_clock).export_json(processor)
    text = path.read_text(encoding="utf-8")
    document = json.loads(text)
    assert text == json.dumps(document, separators=(",", ":"), ensure_ascii=False)
    assert "\n" not in text
    assert "Cien años de soledad" in text
    assert document["metadata"]["record_counts"]["works"] == 4
    assert document["metadata"]["kpis"]["total_works"] == 4
    assert len(document["records"]) == 4


def test_empty_processor_reports(tmp_path: Path) -> None:
    """Reports are still produced when the database is empty."""
    paths = ReportGenerator(tmp_path, clock=fixed_clock).generate_all(DataProcessor([]))
    assert all(path.is_file() for path in paths.values())


def test_report_errors_are_wrapped(processor: DataProcessor, tmp_path: Path) -> None:
    """An unusable output folder raises ``ReportGenerationError``."""
    blocker = tmp_path / "not_a_folder"
    blocker.write_text("x", encoding="utf-8")
    with pytest.raises(ReportGenerationError):
        ReportGenerator(blocker / "outputs").export_json(processor)
