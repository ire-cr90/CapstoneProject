"""Multi-format report generation: styled Excel workbook, CSV summary and minified JSON."""

from __future__ import annotations

import csv
import json
from collections.abc import Callable, Iterable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.config import get_settings
from src.exceptions import ReportGenerationError
from src.logger import get_logger
from src.processing.analytics import DataProcessor

logger = get_logger(__name__)

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
MIN_COLUMN_WIDTH = 8
MAX_COLUMN_WIDTH = 60
TOP_N = 15

COLUMN_FORMATS: dict[str, str] = {
    "year": "0",
    "first_publish_year": "0",
    "decade": "0",
    "count": "#,##0",
    "works": "#,##0",
    "edition_count": "#,##0",
    "pages": "0",
    "works_count": "#,##0",
    "sync_seconds": "0.000",
    "async_seconds": "0.000",
    "speedup": '0.00"x"',
    "run_at": "yyyy-mm-dd hh:mm:ss",
}
SUMMARY_FORMATS: dict[str, str] = {
    "cache_hit_ratio": "0.00%",
    "data_completeness_pct": '0.00"%"',
    "average_async_speedup": '0.00"x"',
    "mean_speedup": '0.00"x"',
    "best_speedup": '0.00"x"',
    "average_edition_count": "#,##0.00",
    "mean_sync_seconds": "0.000",
    "mean_async_seconds": "0.000",
    "generated_at": "yyyy-mm-dd hh:mm:ss",
}
BENCHMARK_COLUMNS: tuple[str, ...] = (
    "run_at",
    "subject",
    "pages",
    "works_count",
    "sync_seconds",
    "async_seconds",
    "speedup",
)


def _cell_value(value: Any) -> Any:
    """Convert pandas and Python values into something openpyxl can store."""
    if isinstance(value, list | tuple | set):
        return ", ".join(str(item) for item in value)
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        timestamp = value.tz_convert(None) if value.tzinfo is not None else value
        return timestamp.to_pydatetime()
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(UTC).replace(tzinfo=None)
    if hasattr(value, "item"):
        return value.item()
    return value


def _label(name: str) -> str:
    """Turn a snake_case column name into a readable header."""
    return name.replace("_", " ").title()


class ReportGenerator:
    """Writes the analytics results to Excel, CSV and JSON files in ``output_dir``."""

    def __init__(
        self,
        output_dir: str | Path | None = None,
        clock: Callable[[], datetime] = datetime.now,
    ) -> None:
        """Remember the output folder and the clock used to date the file names."""
        self.output_dir = Path(output_dir if output_dir is not None else get_settings().output_dir)
        self._clock = clock

    def _prepare_output(self, prefix: str, extension: str) -> Path:
        """Create the output folder if needed and return the dated file path."""
        try:
            self.output_dir.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise ReportGenerationError(
                f"Cannot create output folder {self.output_dir}: {exc}"
            ) from exc
        return self.output_dir / f"{prefix}_{self._clock():%Y%m%d}.{extension}"

    def generate_all(self, processor: DataProcessor) -> dict[str, Path]:
        """Generate the Excel, CSV and JSON reports and return their paths."""
        return {
            "excel": self.export_excel(processor),
            "csv": self.export_csv_summary(processor),
            "json": self.export_json(processor),
        }

    def export_excel(self, processor: DataProcessor) -> Path:
        """Write ``library_report_YYYYMMDD.xlsx`` with summary, data and benchmark sheets."""
        path = self._prepare_output("library_report", "xlsx")
        try:
            workbook = Workbook()
            summary = workbook.active
            summary.title = "Summary"
            self._write_summary(summary, processor)
            self._write_frame(workbook, "Works", processor.works_table())
            years_sheet = self._write_frame(
                workbook, "Top Years", processor.top_publication_years(TOP_N)
            )
            self._add_years_chart(years_sheet)
            self._write_frame(workbook, "Languages", processor.top_languages(TOP_N))
            self._write_frame(workbook, "Authors", processor.top_authors(TOP_N))
            self._write_frame(workbook, "Subjects", processor.top_subjects(TOP_N))
            benchmarks = processor.benchmarks.reindex(columns=list(BENCHMARK_COLUMNS))
            self._write_frame(workbook, "Benchmarks", benchmarks)
            workbook.save(path)
        except (OSError, ValueError, TypeError) as exc:
            raise ReportGenerationError(f"Could not write Excel report {path}: {exc}") from exc
        logger.info("Excel report written to %s", path)
        return path

    def _write_summary(self, sheet: Worksheet, processor: DataProcessor) -> None:
        """Fill the Summary sheet with KPIs, missing values and benchmark metrics."""
        rows: list[tuple[str, str, Any]] = [
            ("KPI", key, value) for key, value in processor.kpis().items()
        ]
        rows += [
            ("Missing values", key, value)
            for key, value in processor.missing_values_report().items()
        ]
        rows += [
            ("Performance", key, value) for key, value in processor.benchmark_summary().items()
        ]
        rows.append(("Report", "generated_at", datetime.now(UTC).replace(tzinfo=None)))
        labelled = [(section, _label(metric), value) for section, metric, value in rows]
        self._write_rows(sheet, ["Section", "Metric", "Value"], labelled)
        for row_index, (_, metric, _) in enumerate(rows, start=2):
            cell = sheet.cell(row=row_index, column=3)
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = SUMMARY_FORMATS.get(metric, cell.number_format)
        self._autosize(sheet)

    def _write_frame(self, workbook: Workbook, title: str, frame: pd.DataFrame) -> Worksheet:
        """Add a sheet with the DataFrame contents, styled headers and number formats."""
        sheet = workbook.create_sheet(title)
        columns = [str(column) for column in frame.columns]
        rows = frame.itertuples(index=False, name=None)
        self._write_rows(sheet, [_label(column) for column in columns], rows)
        for column_index, column in enumerate(columns, start=1):
            number_format = COLUMN_FORMATS.get(column)
            if number_format is None:
                continue
            for row_index in range(2, sheet.max_row + 1):
                sheet.cell(row=row_index, column=column_index).number_format = number_format
        self._autosize(sheet)
        return sheet

    @staticmethod
    def _write_rows(sheet: Worksheet, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
        """Write a styled header row, the data rows, a frozen header and an autofilter."""
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
        for row in rows:
            sheet.append([_cell_value(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    @staticmethod
    def _autosize(sheet: Worksheet) -> None:
        """Set each column width from its longest value, within sensible limits."""
        for column_cells in sheet.columns:
            longest = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=0,
            )
            letter = get_column_letter(column_cells[0].column)
            width = min(max(longest + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)
            sheet.column_dimensions[letter].width = width

    @staticmethod
    def _add_years_chart(sheet: Worksheet) -> None:
        """Add a bar chart of works per publication year next to the Top Years table."""
        if sheet.max_row < 2:
            return
        chart = BarChart()
        chart.title = "Top Publication Years"
        chart.x_axis.title = "Year"
        chart.y_axis.title = "Works"
        chart.legend = None
        data = Reference(sheet, min_col=2, min_row=1, max_row=sheet.max_row)
        categories = Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.width = 18
        chart.height = 9
        sheet.add_chart(chart, "D2")

    def export_csv_summary(self, processor: DataProcessor) -> Path:
        """Write ``library_summary_YYYYMMDD.csv`` with top years and languages."""
        path = self._prepare_output("library_summary", "csv")
        rows: list[tuple[str, Any, int]] = [
            ("publication_year", _cell_value(row.year), int(row.count))
            for row in processor.top_publication_years(TOP_N).itertuples(index=False)
        ]
        rows += [
            ("language", row.language, int(row.count))
            for row in processor.top_languages(TOP_N).itertuples(index=False)
        ]
        try:
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["category", "value", "count"])
                writer.writerows(rows)
        except OSError as exc:
            raise ReportGenerationError(f"Could not write CSV summary {path}: {exc}") from exc
        logger.info("CSV summary written to %s", path)
        return path

    def export_json(self, processor: DataProcessor) -> Path:
        """Write ``library_export_YYYYMMDD.json``: minified metadata and cleaned records."""
        path = self._prepare_output("library_export", "json")
        records = processor.export_records()
        document = {
            "metadata": {
                "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
                "source": "Open Library",
                "record_counts": {
                    "works": len(records),
                    "authors": len(processor.authors),
                    "work_authors": len(processor.work_authors),
                    "benchmarks": len(processor.benchmarks),
                },
                "kpis": processor.kpis(),
                "missing_values": processor.missing_values_report(),
            },
            "records": records,
        }
        try:
            payload = json.dumps(document, separators=(",", ":"), ensure_ascii=False)
            path.write_text(payload, encoding="utf-8")
        except (OSError, TypeError, ValueError) as exc:
            raise ReportGenerationError(f"Could not write JSON export {path}: {exc}") from exc
        logger.info("JSON export written to %s", path)
        return path
